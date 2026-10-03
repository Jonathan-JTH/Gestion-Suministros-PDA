#!/usr/bin/env python3
"""Actualiza el docx de tesis: solo reemplazos puntuales + anexo de aclaraciones; conserva APA y figuras."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
BASE = Path(__file__).resolve().parent
ORIGINAL = BASE / "Gestion-Suministros-Proyecto-G1-ORIGINAL.docx"
OUTPUT = BASE / "Gestion-Suministros-Proyecto-G1-ACTUALIZADO.docx"
LOG = BASE / "COMPARATIVO-CAMBIOS-DOCUMENTO.md"

REPLACEMENTS: list[tuple[str, str]] = [
    (
        "permitiendo realizar solicitudes de suministros sin necesidad de autenticación",
        "permitiendo que cada sucursal registre solicitudes mediante usuarios institucionales autenticados (rol sucursal), vinculados a su sede",
    ),
    (
        "mientras que el personal administrativo contará con una consola de gestión con control de acceso por roles",
        "mientras que el personal de soporte técnico y administración contará con una consola interna con control de acceso por roles y permisos",
    ),
    (
        "registro de solicitudes sin autenticación",
        "registro de solicitudes con autenticación de usuario de sucursal (credenciales asignadas por administración)",
    ),
    (
        "captura estructurada de datos obligatorios: sucursal, ubicación, puesto o área, correo electrónico, modelo del equipo, tipo de solicitud y cantidad requerida",
        "captura de datos: suministro del catálogo, cantidad solicitada y observaciones; la sucursal y el solicitante se obtienen del usuario autenticado",
    ),
    (
        "generación automática de identificadores únicos por tipo de movimiento:",
        "generación de un identificador único secuencial por solicitud (número de registro). En la implementación se utiliza catálogo de suministros y detalle de cantidades (los códigos ENT, DEV y OC del análisis inicial se documentan como evolución del diseño en la nota final)",
    ),
    (
        "envío automático de confirmación al solicitante mediante correo electrónico",
        "envío automático de correo de confirmación al solicitante (correo del usuario de sucursal) y notificación al buzón de soporte técnico (TI) configurado en el sistema",
    ),
    (
        "Portal de solicitudes sin autenticación para usuarios solicitantes.",
        "Módulo de solicitudes para usuarios de sucursal autenticados (rol sucursal), con acceso restringido a registro y consulta de sus propias solicitudes.",
    ),
    (
        "cada suministro estará vinculado a la impresora correspondiente",
        "cada suministro forma parte de un catálogo institucional y se relaciona con el inventario por sucursal",
    ),
    (
        "como usuarios, impresoras, suministros, solicitudes y bitacora_movimientos",
        "como usuarios, roles, permisos, sucursales, suministros, solicitudes, detalle de solicitudes, inventario, movimiento de inventario y bitácora",
    ),
    (
        "Los modelos representan las tablas de la base de datos (usuarios, impresoras, suministros, solicitudes, bitacora_movimientos).",
        "Los modelos representan las tablas de la base de datos (usuarios, roles, permisos, sucursales, suministros, solicitudes, detalle de solicitudes, inventario, movimientos de inventario y bitácora).",
    ),
    (
        "Se implementa autenticación y roles (admin y usuario) para controlar quién puede registrar solicitudes o despachar suministros.",
        "Se implementa autenticación, roles (administrador, soporte técnico y sucursal) y permisos granulares; el segundo factor de autenticación (TOTP) aplica al personal administrativo y de soporte.",
    ),
    (
        "Laravel facilita el envío de correos automáticos cuando se genera una nueva solicitud de tóner o se despacha un suministro, manteniendo a los administradores y usuarios informados en tiempo real.",
        "Laravel facilita el envío de correos automáticos: aviso a TI al registrar una solicitud, confirmación al solicitante, copia obligatoria a TI al atender la solicitud y aviso al correo que soporte indique manualmente al despachar; la activación se administra desde el módulo de configuración y en la tabla parametros_sistema.",
    ),
]

APPENDIX_TITLE = "Nota de alineación con la implementación (Capítulos V y VI)"

APPENDIX_SECTIONS: list[tuple[str, str]] = [
    (
        "Alcance del documento",
        "El Capítulo V (modelo lógico de once entidades, diagrama de colores y casos de uso) se mantiene como diseño aprobado del proyecto. "
        "Las modificaciones de este anexo no sustituyen figuras, MER ni diagramas de secuencia del Capítulo V; solo aclaran cómo se materializó la solución en software.",
    ),
    (
        "Capítulo I y módulo de solicitudes",
        "Se adoptó autenticación por usuario de sucursal (rol sucursal) en lugar de un portal público anónimo, para trazabilidad y control de acceso (RNF-01 a RNF-04).",
    ),
    (
        "Capítulo V — modelo lógico (11 tablas de negocio)",
        "Las once entidades del MER — roles, permisos, rol_permiso, usuarios, sucursales, suministros, inventario, solicitudes, detalle_solicitudes, movimiento_inventario y bitácora — "
        "corresponden a las tablas homónimas en MySQL. Las relaciones y el ciclo de vida de la solicitud (pendiente, en_proceso, atendida, rechazada) se respetan en la implementación.",
    ),
    (
        "Capítulo VI — base de datos física en MySQL Workbench",
        "Además de las once tablas de negocio, la base gestion_suministros incluye tablas auxiliares del framework Laravel (migrations, cache, cache_locks, jobs, job_batches, failed_jobs, sessions, password_reset_tokens) "
        "y la tabla parametros_sistema (configuración operativa de notificaciones por correo). "
        "Las tablas desafios_2fa y codigos_2fa forman parte del diseño de seguridad; la versión desplegada utiliza TOTP (Google Authenticator) con el secreto almacenado en usuarios (totp_secreto). "
        "Estas tablas técnicas no forman parte del MER del Capítulo V y pueden aparecer vacías hasta que el framework o un flujo opcional las utilice.",
    ),
    (
        "Equivalencias de nomenclatura (lógico → físico)",
        "En Laravel convencional: id_usuario/id_rol → id y claves foráneas rol_id, sucursal_id; password_hash → password (hash bcrypt); "
        "mfa_secret → totp_secreto; existencia en inventario → cantidad; estados en minúsculas en ENUM. "
        "La solicitud incluye atributos adicionales de auditoría de despacho y correo (correo_solicitante, correo_destinatario_despacho, nota_despacho, despachado_por_usuario_id) alineados con RF de notificación y trazabilidad.",
    ),
    (
        "Notificaciones por correo electrónico",
        "El envío utiliza SMTP (cuenta remitente configurable en el servidor, por ejemplo correo institucional con contraseña de aplicación). "
        "Los avisos obligatorios a soporte técnico se dirigen al buzón NOTIFICACION_SOPORTE_EMAIL (soporteti@grupofabrigas.com). "
        "Al registrar una solicitud se notifica a TI y se confirma al solicitante; al atender se envía copia a TI y, opcionalmente, un aviso al correo que indique soporte. "
        "Cada intento de envío se registra en la bitácora del sistema (acciones email_enviado / email_fallido).",
    ),
    (
        "Bitácora y logs",
        "La auditoría funcional consultable en la aplicación corresponde a la tabla bitácora (módulo, acción, detalle, usuario, IP). "
        "Los archivos laravel.log en el servidor son registros técnicos del framework y no sustituyen la bitácora de negocio.",
    ),
]


def iter_all_paragraphs(doc: Document):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def replace_in_paragraph(paragraph, old: str, new: str) -> bool:
    full = paragraph.text
    if old not in full:
        return False
    new_text = full.replace(old, new, 1)
    if len(paragraph.runs) == 0:
        paragraph.add_run(new_text)
        return True
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""
    return True


def add_appendix(doc: Document) -> None:
    doc.add_page_break()
    title_p = doc.add_paragraph()
    title_p.add_run(APPENDIX_TITLE).bold = True
    doc.add_paragraph(
        "Las siguientes aclaraciones enlazan el diseño documentado con la implementación Laravel 12 desplegada, "
        "sin modificar el contenido estructural ni las figuras del documento original."
    )
    for heading, body in APPENDIX_SECTIONS:
        h = doc.add_paragraph()
        h.add_run(heading).bold = True
        doc.add_paragraph(body)


def main() -> int:
    if not ORIGINAL.exists():
        print(f"No encontrado: {ORIGINAL}", file=sys.stderr)
        return 1

    doc = Document(str(ORIGINAL))
    log_lines = [
        "# Comparativo de cambios en el documento de tesis",
        "",
        f"**Original:** `{ORIGINAL.name}`  ",
        f"**Actualizado:** `{OUTPUT.name}`  ",
        "",
        "Se aplicaron **solo sustituciones de texto** en párrafos que contenían frases desalineadas con la implementación. "
        "**No se modificaron** el MER de colores (11 tablas), casos de uso ni diagramas del Capítulo V. "
        "Se agregó un **anexo de aclaraciones** (Cap. V y VI) al final.",
        "",
        "## Reemplazos aplicados",
        "",
    ]

    applied = 0
    for old, new in REPLACEMENTS:
        count = 0
        for p in iter_all_paragraphs(doc):
            if replace_in_paragraph(p, old, new):
                count += 1
        if count:
            applied += count
            log_lines.append(f"### ({count} párrafo(s))")
            log_lines.append("")
            log_lines.append("**Antes (fragmento):**")
            log_lines.append(f"> {old[:200]}{'…' if len(old) > 200 else ''}")
            log_lines.append("")
            log_lines.append("**Después (fragmento):**")
            log_lines.append(f"> {new[:300]}{'…' if len(new) > 300 else ''}")
            log_lines.append("")

    add_appendix(doc)

    log_lines.extend(["## Anexo agregado al final", "", f"**Título:** {APPENDIX_TITLE}", ""])
    for heading, body in APPENDIX_SECTIONS:
        log_lines.append(f"### {heading}")
        log_lines.append("")
        log_lines.append(body)
        log_lines.append("")

    doc.save(str(OUTPUT))
    log_lines.append(f"**Total de reemplazos en párrafos:** {applied}")
    log_lines.extend(
        [
            "",
            "## Cómo comparar en Word",
            "",
            "1. Abra `Gestion-Suministros-Proyecto-G1-ORIGINAL.docx` y `Gestion-Suministros-Proyecto-G1-ACTUALIZADO.docx`.",
            "2. Use **Revisar → Comparar** para ver solo los cambios de texto y el anexo final.",
            "3. Entregue o defienda con el ACTUALIZADO; conserve el ORIGINAL como referencia del diseño inicial.",
            "",
        ]
    )
    LOG.write_text("\n".join(log_lines), encoding="utf-8")
    print(f"Guardado: {OUTPUT}")
    print(f"Log: {LOG}")
    print(f"Reemplazos: {applied}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
